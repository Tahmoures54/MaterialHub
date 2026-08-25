import csv, io, os, secrets
from datetime import date, datetime, timedelta
from flask import Blueprint, jsonify, render_template, request, flash, redirect, url_for
from flask_login import login_required, current_user
from werkzeug.utils import secure_filename
from extensions import db
from models import MaterialRequisition, SupplierMaterial, User, AccessLevel, PurchaseOrder
from models_intelligence import (MaterialTrace, MaterialDocument, SupplierScore, RFQ, RFQSupplier,
    Receipt, SupplierInvoice, ThreeWayMatch, MaterialPriceHistory, ScheduleRisk, RFQStatus, MatchStatus, DocumentType)

intelligence_bp=Blueprint('intelligence',__name__)

def company_filter(q, model):
    return q if current_user.is_admin else q.filter(model.company_name==current_user.company_name)

def score_supplier(supplier_id):
    orders=PurchaseOrder.query.filter_by(supplier_id=supplier_id).all()
    if not orders: return 0
    delivered=sum(1 for o in orders if o.delivered_date and o.issued_date and o.delivered_date<=o.issued_date+timedelta(days=30))
    quality=100
    try:
        from models import QualityControl, InspectionStatus
        qcs=QualityControl.query.filter_by(user_id=supplier_id).all()
        if qcs: quality=100*sum(1 for q in qcs if q.status==InspectionStatus.passed)/len(qcs)
    except Exception: pass
    otif=100*delivered/len(orders)
    score=round(0.45*otif+0.35*quality+0.20*100,1)
    s=SupplierScore.query.filter_by(supplier_id=supplier_id,period=date.today().strftime('%Y-%m')).first()
    if not s:
        s=SupplierScore(supplier_id=supplier_id,period=date.today().strftime('%Y-%m'),company_name=current_user.company_name)
        db.session.add(s)
    s.quality_score=round(quality,1); s.otif_score=round(otif,1); s.price_score=100; s.responsiveness_score=100; s.lead_time_score=100; s.overall_score=score; s.orders_count=len(orders)
    db.session.commit(); return score

def readiness(item_code, required_qty=0, required_date=None):
    stock=sum((x.received_qty or 0) for x in __import__('models').WarehouseInventory.query.filter_by(item_code=item_code).all())
    open_po=0
    for po in PurchaseOrder.query.filter(PurchaseOrder.status.in_([__import__('models').PurchaseOrderStatus.pending,__import__('models').PurchaseOrderStatus.issued])).all():
        for it in po.items:
            if it.material_requisition and it.material_requisition.item_code==item_code: open_po+=it.quantity or 0
    coverage=stock+open_po
    qty_score=min(100, (coverage/required_qty*100) if required_qty else 100)
    date_score=100
    if required_date:
        days=(required_date-date.today()).days
        if days<0: date_score=0
        elif days<14 and open_po==0: date_score=30
        elif days<30 and open_po==0: date_score=60
    return round(0.7*qty_score+0.3*date_score,1),stock,open_po

def generate_copilot():
    alerts=[]
    for mr in company_filter(MaterialRequisition.query,MaterialRequisition).filter(MaterialRequisition.status!='rejected').order_by(MaterialRequisition.required_date.asc()).limit(100).all():
        idx,stock,open_po=readiness(mr.item_code,mr.quantity,mr.required_date)
        if idx<75:
            days=max((mr.required_date-date.today()).days,0) if mr.required_date else 999
            alerts.append({'item_code':mr.item_code,'description':mr.material_description,'readiness':idx,'days_to_required':days,'stock':stock,'open_po':open_po,
             'message':f"{mr.item_code} has {days} days until required date and only {stock:g} in stock with {open_po:g} on open PO.",
             'action':'Expedite open PO' if open_po else 'Create RFQ and activate qualified alternate supplier'})
    return sorted(alerts,key=lambda x:x['readiness'])[:12]

@intelligence_bp.route('/material-intelligence')
@login_required
def dashboard():
    traces=company_filter(MaterialTrace.query,MaterialTrace).order_by(MaterialTrace.created_at.desc()).limit(10).all()
    docs=company_filter(MaterialDocument.query,MaterialDocument).filter_by(approved=False).limit(10).all()
    risks=generate_copilot()
    scores=company_filter(SupplierScore.query,SupplierScore).order_by(SupplierScore.overall_score.desc()).limit(8).all()
    matches=company_filter(ThreeWayMatch.query,ThreeWayMatch).filter(ThreeWayMatch.status==MatchStatus.exception).limit(10).all()
    return render_template('intelligence/material_intelligence.html',traces=traces,docs=docs,risks=risks,scores=scores,matches=matches)

@intelligence_bp.route('/api/material-intelligence')
@login_required
def api_dashboard():
    return jsonify({'copilot':generate_copilot(),'supplier_scores':[{'supplier_id':s.supplier_id,'score':s.overall_score,'otif':s.otif_score} for s in company_filter(SupplierScore.query,SupplierScore).order_by(SupplierScore.overall_score.desc()).limit(20).all()]})

@intelligence_bp.route('/traceability',methods=['GET','POST'])
@login_required
def traceability():
    if request.method=='POST':
        t=MaterialTrace(trace_code=request.form['trace_code'],item_code=request.form['item_code'],material_description=request.form['material_description'],heat_no=request.form.get('heat_no'),lot_no=request.form.get('lot_no'),serial_no=request.form.get('serial_no'),batch_no=request.form.get('batch_no'),quantity=float(request.form.get('quantity') or 0),unit=request.form.get('unit'),po_no=request.form.get('po_no'),delivery_id=request.form.get('delivery_id'),warehouse_id=request.form.get('warehouse_id'),supplier_name=request.form.get('supplier_name'),project_no=request.form.get('project_no'),company_name=current_user.company_name)
        db.session.add(t); db.session.commit(); flash('Traceability record created.','success')
    return render_template('intelligence/traceability.html',traces=company_filter(MaterialTrace.query,MaterialTrace).order_by(MaterialTrace.created_at.desc()).limit(100).all())

@intelligence_bp.route('/api/trace/<trace_code>')
@login_required
def trace_api(trace_code):
    t=company_filter(MaterialTrace.query,MaterialTrace).filter_by(trace_code=trace_code).first_or_404()
    return jsonify({'trace_code':t.trace_code,'item_code':t.item_code,'heat_no':t.heat_no,'lot_no':t.lot_no,'serial_no':t.serial_no,'batch_no':t.batch_no,'po_no':t.po_no,'delivery_id':t.delivery_id,'warehouse_id':t.warehouse_id,'supplier':t.supplier_name,'project':t.project_no,'documents':[{'type':d.document_type.value,'file':d.file_name,'approved':d.approved} for d in t.documents]})

@intelligence_bp.route('/rfq/auto/<int:mr_id>',methods=['POST'])
@login_required
def auto_rfq(mr_id):
    mr=company_filter(MaterialRequisition.query,MaterialRequisition).filter_by(id=mr_id).first_or_404()
    rfq=RFQ(rfq_no='RFQ-'+date.today().strftime('%Y%m%d')+'-'+secrets.token_hex(2).upper(),mr_id=mr.id,status=RFQStatus.sent,due_date=date.today()+timedelta(days=5),target_quantity=mr.quantity,unit=mr.unit_of_measure,specification=f'{mr.material_description} | {mr.standard_specification or "Open specification"}',company_name=current_user.company_name,created_by=current_user.id)
    db.session.add(rfq); db.session.flush()
    suppliers=User.query.filter(User.access_level==AccessLevel.supplier,User.company_name==current_user.company_name).limit(20).all()
    if not suppliers: suppliers=User.query.filter(User.access_level==AccessLevel.supplier).limit(20).all()
    for s in suppliers: db.session.add(RFQSupplier(rfq_id=rfq.id,supplier_id=s.id))
    db.session.commit(); return jsonify({'rfq_no':rfq.rfq_no,'invited_suppliers':len(suppliers),'status':rfq.status.value})

@intelligence_bp.route('/rfq/compare/<int:rfq_id>')
@login_required
def compare_rfq(rfq_id):
    rfq=RFQ.query.get_or_404(rfq_id)
    quotes=RFQSupplier.query.filter_by(rfq_id=rfq.id).order_by(RFQSupplier.total_score.desc(),RFQSupplier.quoted_price.asc()).all()
    return jsonify({'rfq_no':rfq.rfq_no,'comparison':[{'supplier':q.supplier.company_name,'price':q.quoted_price,'lead_time_days':q.lead_time_days,'quality_score':q.quality_score,'total_score':q.total_score,'status':q.status} for q in quotes]})

@intelligence_bp.route('/three-way-match/<int:po_id>',methods=['GET','POST'])
@login_required
def three_way_match(po_id):
    po=PurchaseOrder.query.get_or_404(po_id); receipt=Receipt.query.filter_by(po_id=po.id).order_by(Receipt.id.desc()).first(); invoice=SupplierInvoice.query.filter_by(po_id=po.id).order_by(SupplierInvoice.id.desc()).first()
    if request.method=='POST':
        receipt=Receipt(receipt_no=request.form['receipt_no'],po_id=po.id,received_qty=float(request.form.get('received_qty') or 0),accepted_qty=float(request.form.get('accepted_qty') or 0),receipt_date=date.today(),company_name=current_user.company_name) if not receipt else receipt
        if not receipt.id: db.session.add(receipt); db.session.flush()
        invoice=SupplierInvoice(invoice_no=request.form['invoice_no'],po_id=po.id,invoice_amount=float(request.form.get('invoice_amount') or 0),invoice_date=date.today(),company_name=current_user.company_name) if not invoice else invoice
        if not invoice.id: db.session.add(invoice); db.session.flush()
        variance=abs((po.total_price or 0)-invoice.invoice_amount)
        status=MatchStatus.matched if variance <= max(1,po.total_price*.02) and receipt.accepted_qty>=0 else MatchStatus.exception
        m=ThreeWayMatch(po_id=po.id,receipt_id=receipt.id,invoice_id=invoice.id,po_amount=po.total_price or 0,receipt_amount=receipt.accepted_qty or 0,invoice_amount=invoice.invoice_amount,variance_amount=variance,status=status,notes='Automatic 3-way match')
        db.session.add(m); db.session.commit(); return jsonify({'status':status.value,'variance':variance})
    return jsonify({'po':po.order_no,'po_amount':po.total_price,'receipt':receipt.accepted_qty if receipt else None,'invoice':invoice.invoice_amount if invoice else None})

@intelligence_bp.route('/excel-import',methods=['GET','POST'])
@login_required
def excel_import():
    if request.method=='POST':
        f=request.files.get('file')
        if not f: flash('Select a CSV or XLSX file.','danger'); return redirect(request.url)
        rows=[]
        try:
            if f.filename.lower().endswith('.csv'):
                rows=list(csv.DictReader(io.StringIO(f.stream.read().decode('utf-8-sig'))))
            else:
                import openpyxl
                wb=openpyxl.load_workbook(f,read_only=True,data_only=True); ws=wb.active; headers=[str(c.value).strip() if c.value is not None else '' for c in next(ws.iter_rows())]
                rows=[dict(zip(headers,[c.value for c in row])) for row in ws.iter_rows()]
            errors=[]; valid=[]
            required={'item_code','material_description','quantity','required_date'}
            for n,r in enumerate(rows,2):
                missing=[x for x in required if not r.get(x)]
                if missing: errors.append({'row':n,'error':'Missing: '+', '.join(missing)})
                else: valid.append(r)
            return render_template('intelligence/excel_import.html',rows=valid[:100],errors=errors)
        except Exception as e: flash(f'Import analysis failed: {e}','danger')
    return render_template('intelligence/excel_import.html',rows=[],errors=[])

@intelligence_bp.route('/api/qr/<trace_code>')
@login_required
def qr_code(trace_code):
    import qrcode
    from flask import Response
    t=company_filter(MaterialTrace.query,MaterialTrace).filter_by(trace_code=trace_code).first_or_404()
    img=qrcode.make(t.trace_code)
    buf=io.BytesIO(); img.save(buf,format='PNG'); buf.seek(0)
    return Response(buf.getvalue(),mimetype='image/png')

@intelligence_bp.route('/supplier-score/<int:supplier_id>',methods=['POST'])
@login_required
def supplier_score(supplier_id):
    score=score_supplier(supplier_id)
    return jsonify({'supplier_id':supplier_id,'overall_score':score})

@intelligence_bp.route('/price-history/<item_code>')
@login_required
def price_history(item_code):
    rows=company_filter(MaterialPriceHistory.query,MaterialPriceHistory).filter_by(item_code=item_code).order_by(MaterialPriceHistory.recorded_at.desc()).limit(50).all()
    return jsonify([{'price':r.price,'currency':r.currency,'unit':r.unit,'lead_time_days':r.lead_time_days,'recorded_at':r.recorded_at.isoformat(),'supplier_id':r.supplier_id} for r in rows])

@intelligence_bp.route('/lead-time/<item_code>')
@login_required
def lead_time(item_code):
    rows=company_filter(MaterialPriceHistory.query,MaterialPriceHistory).filter_by(item_code=item_code).all()
    values=[r.lead_time_days for r in rows if r.lead_time_days is not None]
    avg=round(sum(values)/len(values),1) if values else None
    return jsonify({'item_code':item_code,'observations':len(values),'average_lead_time_days':avg,'best_lead_time_days':min(values) if values else None,'worst_lead_time_days':max(values) if values else None})
